"""
Officer Bulk Import & Google Sheets Template Service for DPIG.
Allows SuperAdmins and Departmental Admins to download standardized Google Sheet / Excel templates
with built-in dropdown DataValidation, bulk-provision officers across all 7 workflow escalation levels
(DL1 -> DL2 -> DL3 -> CL1 -> CL2 -> CL3 -> CM_OFFICE), departments, and zones, and export active rosters.
"""
import io
import re
import logging
from typing import Dict, Any, List, Tuple, Optional
import pandas as pd
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from django.db import transaction
from django.contrib.auth.models import User
from core.models import (
    Cluster, Zone, Ward, Department, UserProfile,
    WORKFLOW_LEVEL_CODES, LEVEL_CODE_TO_INT, LEVEL_INT_TO_CODE, LEVEL_LABELS
)

logger = logging.getLogger(__name__)

OFFICER_TEMPLATE_COLUMNS = [
    'Username',
    'Full_Name',
    'Email',
    'Phone',
    'Role',
    'Workflow_Level',
    'Zone_Number',
    'Ward_Number',
    'Department_Code',
    'Designation',
    'Password',
]

SAMPLE_OFFICERS_DATA = [
    {
        'Username': 'dl1_ae_roads_z5',
        'Full_Name': 'K. Ramanathan AE',
        'Email': 'dl1.roads.z5@chennaicorporation.gov.in',
        'Phone': '9840011101',
        'Role': 'WARD_OFFICER',
        'Workflow_Level': 'DL1',
        'Zone_Number': 5,
        'Ward_Number': 52,
        'Department_Code': 'ROADS',
        'Designation': 'Assistant Engineer - Roads (DL1)',
        'Password': 'Admin@Dpig2026',
    },
    {
        'Username': 'dl1_si_swm_z5',
        'Full_Name': 'S. Lakshmi SI',
        'Email': 'dl1.swm.z5@chennaicorporation.gov.in',
        'Phone': '9840011102',
        'Role': 'FIELD_STAFF',
        'Workflow_Level': 'DL1',
        'Zone_Number': 5,
        'Ward_Number': 52,
        'Department_Code': 'SWM',
        'Designation': 'Sanitary Inspector - SWM (DL1)',
        'Password': 'Admin@Dpig2026',
    },
    {
        'Username': 'dl2_aee_roads_z5',
        'Full_Name': 'M. Vignesh AEE',
        'Email': 'dl2.roads.z5@chennaicorporation.gov.in',
        'Phone': '9840011103',
        'Role': 'WARD_OFFICER',
        'Workflow_Level': 'DL2',
        'Zone_Number': 5,
        'Ward_Number': '',
        'Department_Code': 'ROADS',
        'Designation': 'Assistant Executive Engineer - Zone 5 (DL2)',
        'Password': 'Admin@Dpig2026',
    },
    {
        'Username': 'dl3_ee_zo5',
        'Full_Name': 'Dr. M. Suresh ZO',
        'Email': 'zo5@chennaicorp.gov.in',
        'Phone': '9840011105',
        'Role': 'ZONAL_OFFICER',
        'Workflow_Level': 'DL3',
        'Zone_Number': 5,
        'Ward_Number': '',
        'Department_Code': '',
        'Designation': 'Zonal Officer - Royapuram (DL3)',
        'Password': 'Admin@Dpig2026',
    },
    {
        'Username': 'cl1_se_roads',
        'Full_Name': 'Tmt. R. Sumathi SE',
        'Email': 'cl1.roads@chennaicorporation.gov.in',
        'Phone': '9840011120',
        'Role': 'COMMISSIONER',
        'Workflow_Level': 'CL1',
        'Zone_Number': '',
        'Ward_Number': '',
        'Department_Code': 'ROADS',
        'Designation': 'Superintending Engineer - Citywide Roads (CL1)',
        'Password': 'Admin@Dpig2026',
    },
    {
        'Username': 'cl2_dc_works',
        'Full_Name': 'Thiru. S. Rajendran IAS',
        'Email': 'dc.works@chennaicorporation.gov.in',
        'Phone': '9840011125',
        'Role': 'COMMISSIONER',
        'Workflow_Level': 'CL2',
        'Zone_Number': '',
        'Ward_Number': '',
        'Department_Code': '',
        'Designation': 'Deputy Commissioner - Works (CL2)',
        'Password': 'Admin@Dpig2026',
    },
    {
        'Username': 'cl3_commissioner',
        'Full_Name': 'J. Kumaragurubaran IAS',
        'Email': 'commissioner@chennaicorporation.gov.in',
        'Phone': '9840011111',
        'Role': 'COMMISSIONER',
        'Workflow_Level': 'CL3',
        'Zone_Number': '',
        'Ward_Number': '',
        'Department_Code': '',
        'Designation': 'Principal Secretary / Commissioner (CL3)',
        'Password': 'Admin@Dpig2026',
    },
    {
        'Username': 'cm_special_cell',
        'Full_Name': 'Thiru. K. Annamalai IAS',
        'Email': 'cmcell@tn.gov.in',
        'Phone': '9840099999',
        'Role': 'CM_OFFICE',
        'Workflow_Level': 'CM_OFFICE',
        'Zone_Number': '',
        'Ward_Number': '',
        'Department_Code': '',
        'Designation': 'Special Officer - Chief Minister Executive Cell',
        'Password': 'Admin@Dpig2026',
    },
]


def generate_officer_template_dataframe(department=None) -> pd.DataFrame:
    """Creates a sample DataFrame matching the Google Sheet officer bulk upload schema."""
    data = [dict(row) for row in SAMPLE_OFFICERS_DATA]
    if department:
        dept_code = getattr(department, 'code', str(department)).upper()
        for row in data:
            if row.get('Department_Code') or row.get('Workflow_Level') in ['DL1', 'DL2', 'CL1']:
                row['Department_Code'] = dept_code
    return pd.DataFrame(data, columns=OFFICER_TEMPLATE_COLUMNS)


def _style_officer_workbook(wb: openpyxl.Workbook, cluster=None):
    """Applies government blue header styling, column widths, and Excel dropdown validations."""
    ws = wb['Officers_Bulk_Upload']
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 26

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    # 1. Role DataValidation (Column E)
    roles = [r[0] for r in UserProfile.ROLES if r[0] != 'CITIZEN']
    dv_role = DataValidation(
        type="list",
        formula1=f'"{",".join(roles)}"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Invalid Role",
        error="Please select a valid government officer role from the dropdown list."
    )
    ws.add_data_validation(dv_role)
    dv_role.add("E2:E500")

    # 2. Workflow Level DataValidation (Column F)
    dv_wf = DataValidation(
        type="list",
        formula1=f'"{",".join(WORKFLOW_LEVEL_CODES)}"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Invalid Workflow Level",
        error="Please select a valid Workflow Escalation Level (DL1 to CM_OFFICE)."
    )
    ws.add_data_validation(dv_wf)
    dv_wf.add("F2:F500")

    # 3. Department Code DataValidation (Column I)
    dept_qs = Department.objects.filter(cluster=cluster) if cluster else Department.objects.all()
    dept_codes = list(dept_qs.values_list('code', flat=True).distinct())
    if not dept_codes:
        dept_codes = ['ROADS', 'SWM', 'SWD', 'LIGHTS', 'HEALTH', 'PARKS', 'REVENUE', 'WATER']
    dept_str = ",".join(sorted(dept_codes))
    dv_dept = DataValidation(
        type="list",
        formula1=f'"{dept_str}"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Invalid Department Code",
        error="Please select an approved Department Code."
    )
    ws.add_data_validation(dv_dept)
    dv_dept.add("I2:I500")


def export_officer_template(file_format: str = 'xlsx', department=None, cluster=None) -> Tuple[bytes, str, str]:
    """
    Renders the Google Sheet template as XLSX (with native Excel dropdown validation) or CSV bytes.
    Optionally customized for a specific department.
    Returns (bytes_data, content_type, filename).
    """
    df = generate_officer_template_dataframe(department=department)

    dept_slug = ""
    if department:
        d_code = getattr(department, 'code', str(department)).lower()
        dept_slug = f"_{d_code}"

    if file_format.lower() == 'csv':
        csv_buffer = io.StringIO()
        df.to_csv(csv_buffer, index=False)
        return (
            csv_buffer.getvalue().encode('utf-8'),
            'text/csv; charset=utf-8',
            f'dpig_officer_bulk_upload_template{dept_slug}.csv'
        )

    # Excel .xlsx with dropdown data validations and Reference sheet
    excel_buffer = io.BytesIO()
    with pd.ExcelWriter(excel_buffer, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Officers_Bulk_Upload', index=False)

        # Instructions & Reference sheet
        active_cluster = cluster or (department.cluster if department else None) or Cluster.objects.filter(is_active=True).first()
        dept_qs = Department.objects.filter(cluster=active_cluster) if active_cluster else Department.objects.all()
        dept_names = ", ".join(f"{d.code} ({d.name})" for d in dept_qs) or "ROADS, SWM, SWD, LIGHTS, HEALTH, PARKS, REVENUE, WATER"

        ref_data = {
            'Field': [
                'Workflow_Level',
                'Roles Supported',
                'Departments Available',
                'Zones Available',
                'Ward Numbers',
                'Default Password'
            ],
            'Allowed Values / Instructions': [
                'DL1 (Dept Level 1), DL2 (Dept Level 2), DL3 (Dept Level 3), CL1 (City Level 1), CL2 (City Level 2), CL3 (City Level 3), CM_OFFICE (Apex)',
                'WARD_OFFICER, FIELD_STAFF, ZONAL_OFFICER, COMMISSIONER, CM_OFFICE, POLICYMAKER, AUDITOR, SUPERADMIN, DEPT_ADMIN',
                dept_names,
                '1 to 15 (Leave blank for Citywide / Statewide officers)',
                '1 to 200 (Optional, for ward-specific officers)',
                'e.g. Admin@Dpig2026 (Optional, defaults to Admin@Dpig2026)'
            ]
        }
        pd.DataFrame(ref_data).to_excel(writer, sheet_name='Instructions_&_Reference', index=False)

    # Reopen with openpyxl to apply styling & DataValidation
    excel_buffer.seek(0)
    wb = openpyxl.load_workbook(excel_buffer)
    _style_officer_workbook(wb, cluster=active_cluster)

    # Style Reference Sheet as well
    ws_ref = wb['Instructions_&_Reference']
    ws_ref.column_dimensions['A'].width = 25
    ws_ref.column_dimensions['B'].width = 90
    for cell in ws_ref[1]:
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")

    out_buffer = io.BytesIO()
    wb.save(out_buffer)

    return (
        out_buffer.getvalue(),
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        f'dpig_officer_bulk_upload_template{dept_slug}.xlsx'
    )


def export_officers_to_excel(cluster=None, department=None, file_format: str = 'xlsx') -> Tuple[bytes, str, str]:
    """
    Exports active officer directory to an Excel workbook or CSV in template-compatible format.
    Allows administrators to backup or modify rosters offline and re-upload.
    """
    profiles = UserProfile.objects.select_related('user', 'zone', 'ward', 'department', 'cluster')
    if cluster:
        profiles = profiles.filter(cluster=cluster)
    if department:
        profiles = profiles.filter(department=department)

    profiles = profiles.exclude(role='CITIZEN').order_by('department__code', 'hierarchy_level', 'user__username')

    rows = []
    for p in profiles:
        wf_code = LEVEL_INT_TO_CODE.get(p.hierarchy_level, 'DL1')
        full_n = p.user.get_full_name() or p.user.username
        rows.append({
            'Username': p.user.username,
            'Full_Name': full_n,
            'Email': p.user.email,
            'Phone': p.phone,
            'Role': p.role,
            'Workflow_Level': wf_code,
            'Zone_Number': p.zone.number if p.zone else '',
            'Ward_Number': p.ward.number if p.ward else '',
            'Department_Code': p.department.code if p.department else '',
            'Designation': p.designation,
            'Password': '',  # left blank for security; re-import retains existing password if blank
        })

    df = pd.DataFrame(rows, columns=OFFICER_TEMPLATE_COLUMNS)
    suffix = f"_{department.code.lower()}" if department else ""

    if file_format.lower() == 'csv':
        csv_buffer = io.StringIO()
        df.to_csv(csv_buffer, index=False)
        return (
            csv_buffer.getvalue().encode('utf-8'),
            'text/csv; charset=utf-8',
            f'dpig_active_officers_roster{suffix}.csv'
        )

    excel_buffer = io.BytesIO()
    with pd.ExcelWriter(excel_buffer, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Officers_Bulk_Upload', index=False)

    excel_buffer.seek(0)
    wb = openpyxl.load_workbook(excel_buffer)
    _style_officer_workbook(wb, cluster=cluster)

    out_buffer = io.BytesIO()
    wb.save(out_buffer)

    return (
        out_buffer.getvalue(),
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        f'dpig_active_officers_roster{suffix}.xlsx'
    )


def parse_and_provision_officers(file_obj, cluster=None, update_existing: bool = True, enforce_department=None) -> Dict[str, Any]:
    """
    Parses an uploaded Google Sheet (CSV or Excel) and provisions/updates officers.
    Optionally enforces a specific Department (e.g. when uploaded by Departmental Admin).
    Returns:
        {
            'status': 'SUCCESS' | 'PARTIAL' | 'ERROR',
            'total_rows': int,
            'created_count': int,
            'updated_count': int,
            'error_count': int,
            'errors': List[str],
            'warnings': List[str],
            'provisioned_users': List[str]
        }
    """
    if cluster is None:
        cluster = Cluster.objects.filter(is_active=True).first()

    email_domain = "chennaicorporation.gov.in"
    if cluster and cluster.contact_email and '@' in cluster.contact_email:
        email_domain = cluster.contact_email.split('@')[1].strip()

    if isinstance(file_obj, pd.DataFrame):
        df = file_obj.copy()
    else:
        filename = getattr(file_obj, 'name', '').lower()
        try:
            if filename.endswith('.csv'):
                df = pd.read_csv(file_obj, dtype=str).fillna('')
            else:
                try:
                    # Smart sheet selection: find sheet with 'officer', 'staff', 'user', or first non-instructions sheet
                    xls = pd.ExcelFile(file_obj)
                    target_sheet = None
                    for s in xls.sheet_names:
                        s_lower = s.lower()
                        if any(kw in s_lower for kw in ('officer', 'staff', 'user', 'roster', 'upload')):
                            target_sheet = s
                            break
                    if not target_sheet:
                        non_inst = [s for s in xls.sheet_names if 'instruction' not in s.lower() and 'reference' not in s.lower()]
                        target_sheet = non_inst[0] if non_inst else xls.sheet_names[0]

                    df = pd.read_excel(xls, sheet_name=target_sheet, dtype=str).fillna('')
                except Exception:
                    if hasattr(file_obj, 'seek'):
                        file_obj.seek(0)
                    df = pd.read_csv(file_obj, dtype=str).fillna('')
        except Exception as e:
            logger.error("Failed to read officer spreadsheet: %s", e)
            return {
                'status': 'ERROR',
                'total_rows': 0,
                'created_count': 0,
                'updated_count': 0,
                'error_count': 1,
                'errors': [f"Could not parse uploaded spreadsheet: {str(e)}"],
                'warnings': [],
                'provisioned_users': []
            }

    if df.empty:
        return {
            'status': 'ERROR',
            'total_rows': 0,
            'created_count': 0,
            'updated_count': 0,
            'error_count': 1,
            'errors': ["The uploaded sheet contains no data rows."],
            'warnings': [],
            'provisioned_users': []
        }

    # Normalize column names: strip, lowercase, replace spaces/dashes with underscores
    col_map = {}
    for c in df.columns:
        norm = re.sub(r'[^a-zA-Z0-9]', '_', str(c).strip().lower())
        col_map[c] = norm
    df = df.rename(columns=col_map)
    df = df.fillna('')

    def clean_cell(val):
        if val is None:
            return ''
        s = str(val).strip()
        return '' if s.lower() in ['nan', 'none', 'null'] else s

    created_count = 0
    updated_count = 0
    errors: List[str] = []
    warnings: List[str] = []
    provisioned_users: List[str] = []

    for idx, row in df.iterrows():
        row_num = idx + 2  # 1-indexed header + data
        username = clean_cell(row.get('username', ''))
        full_name = clean_cell(row.get('full_name', ''))
        email = clean_cell(row.get('email', ''))
        phone = clean_cell(row.get('phone', ''))
        role_input = clean_cell(row.get('role', '')).upper()
        wf_level_input = clean_cell(row.get('workflow_level', '')).upper()
        zone_val = clean_cell(row.get('zone_number', ''))
        ward_val = clean_cell(row.get('ward_number', ''))
        dept_val = clean_cell(row.get('department_code', '')).upper()
        designation = clean_cell(row.get('designation', ''))
        password = clean_cell(row.get('password', ''))

        # Check required fields
        if not username:
            errors.append(f"Row {row_num}: Missing required 'Username'. Skipped.")
            continue

        username = re.sub(r'[^a-zA-Z0-9_.]', '_', username).lower()

        if not full_name:
            full_name = username.replace('_', ' ').replace('.', ' ').title()

        first_name = full_name.split()[0] if full_name else username
        last_name = " ".join(full_name.split()[1:]) if len(full_name.split()) > 1 else ""

        # Resolve Workflow Level & Hierarchy Level
        wf_level = 'DL1'
        hierarchy_level = 1
        if wf_level_input in WORKFLOW_LEVEL_CODES:
            wf_level = wf_level_input
            hierarchy_level = LEVEL_CODE_TO_INT[wf_level]
        elif wf_level_input.isdigit() and int(wf_level_input) in LEVEL_INT_TO_CODE:
            hierarchy_level = int(wf_level_input)
            wf_level = LEVEL_INT_TO_CODE[hierarchy_level]
        elif wf_level_input:
            errors.append(f"Row {row_num}: Invalid Workflow Level '{wf_level_input}'. Expected DL1, DL2, DL3, CL1, CL2, CL3, or CM_OFFICE.")
            continue

        # Resolve Role
        allowed_roles = [
            'WARD_OFFICER', 'FIELD_STAFF', 'ZONAL_OFFICER',
            'COMMISSIONER', 'CM_OFFICE', 'POLICYMAKER', 'AUDITOR', 'SUPERADMIN', 'DEPT_ADMIN'
        ]
        role = role_input
        if not role or role not in allowed_roles:
            if wf_level in ('DL1', 'DL2'):
                role = 'WARD_OFFICER'
            elif wf_level == 'DL3':
                role = 'ZONAL_OFFICER'
            elif wf_level in ('CL1', 'CL2', 'CL3'):
                role = 'COMMISSIONER'
            elif wf_level == 'CM_OFFICE':
                role = 'CM_OFFICE'
            else:
                role = 'WARD_OFFICER'

        # Resolve Zone
        zone_obj = None
        if zone_val:
            try:
                z_int = int(float(zone_val))
                zone_obj = Zone.objects.filter(cluster=cluster, number=z_int).first()
                if not zone_obj:
                    zone_obj = Zone.objects.filter(number=z_int).first()
            except ValueError:
                zone_obj = Zone.objects.filter(name__icontains=zone_val).first()
            if not zone_obj:
                warnings.append(f"Row {row_num} ({username}): Zone '{zone_val}' not found; provisioned without zone.")

        # Resolve Ward
        ward_obj = None
        if ward_val:
            try:
                w_int = int(float(ward_val))
                if zone_obj:
                    ward_obj = Ward.objects.filter(zone=zone_obj, number=w_int).first()
                if not ward_obj:
                    ward_obj = Ward.objects.filter(number=w_int).first()
            except ValueError:
                pass
            if not ward_obj:
                warnings.append(f"Row {row_num} ({username}): Ward '{ward_val}' not found; provisioned without ward.")

        # Resolve Department (enforce_department takes precedence if provided)
        dept_obj = enforce_department
        if not dept_obj and dept_val:
            dept_obj = Department.objects.filter(cluster=cluster, code__iexact=dept_val).first()
            if not dept_obj:
                dept_obj = Department.objects.filter(code__iexact=dept_val).first()
            if not dept_obj:
                warnings.append(f"Row {row_num} ({username}): Department '{dept_val}' not found; provisioned without department.")

        # Default designation if not given
        if not designation:
            designation = f"{LEVEL_LABELS.get(wf_level, wf_level)} - {role.replace('_', ' ').title()}"

        # Create or update User & UserProfile inside atomic block
        try:
            with transaction.atomic():
                user = User.objects.filter(username=username).first()
                is_create = user is None

                if is_create:
                    assigned_pwd = password or 'Admin@Dpig2026'
                    user = User.objects.create_user(
                        username=username,
                        email=email or f"{username}@{email_domain}",
                        password=assigned_pwd,
                        first_name=first_name,
                        last_name=last_name
                    )
                    user.is_staff = True
                    user.is_superuser = (role == 'SUPERADMIN')
                    user.save()
                    created_count += 1
                elif update_existing:
                    user.first_name = first_name
                    user.last_name = last_name
                    if email:
                        user.email = email
                    if password:
                        user.set_password(password)
                    user.is_staff = True
                    user.is_superuser = (role == 'SUPERADMIN')
                    user.save()
                    updated_count += 1
                else:
                    errors.append(f"Row {row_num}: User '{username}' exists and overwrite was disabled.")
                    continue

                profile, _ = UserProfile.objects.update_or_create(
                    user=user,
                    defaults={
                        'cluster': cluster,
                        'role': role,
                        'hierarchy_level': hierarchy_level,
                        'zone': zone_obj,
                        'ward': ward_obj,
                        'department': dept_obj,
                        'designation': designation,
                        'phone': phone
                    }
                )
                provisioned_users.append(f"{username} ({wf_level})")

        except Exception as e:
            logger.error("Error provisioning officer at row %d (%s): %s", row_num, username, e)
            errors.append(f"Row {row_num} ({username}): {str(e)}")

    total_rows = len(df)
    status = 'SUCCESS' if not errors else ('PARTIAL' if (created_count or updated_count) else 'ERROR')

    return {
        'status': status,
        'total_rows': total_rows,
        'created_count': created_count,
        'updated_count': updated_count,
        'error_count': len(errors),
        'errors': errors,
        'warnings': warnings,
        'provisioned_users': provisioned_users
    }
